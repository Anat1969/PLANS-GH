"""יצוא טבלת תוכניות הבינוי לאקסל.

מקורות: RAW_DATA ב-index.html (רשימת התוכניות) + seed-data.json (פרטי יזם/אדריכל,
תאריכי ועדה והחלטות). אינו כולל עריכות שנשמרו רק בענן (Supabase) או בדפדפן.
הרצה: python3 tools/export-plans-xlsx.py  →  exports/plans.xlsx
"""
import json, re, datetime, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
html = (ROOT / 'index.html').read_text(encoding='utf-8')
raw = re.search(r'const RAW_DATA = (\[.*?\n\]);', html, re.S).group(1)
raw = re.sub(r',\s*\]$', ']', raw.strip())
RAW = json.loads(raw)
seed = json.loads((ROOT / 'seed-data.json').read_text(encoding='utf-8'))
F, O = seed['fields'], seed['opinions']

def d(s):
    try: return datetime.date.fromisoformat(s) if s else None
    except ValueError: return None

rows = []
for r in RAW:
    tik, district, planNum, change, submitter, location, purpose, date, handler, status, year, month = r
    # ה-seed ממוספר לפי מס׳ התכנית (תב/XXXX), לא לפי מס׳ התיק
    key = re.sub(r'\D', '', planNum)
    f, o = F.get(key, {}), O.get(key, {})
    opened, committee = d(date), d(o.get('committeeDate'))
    notes = []
    if committee and opened and committee < opened:
        # מסמך ועדה שקודם לפתיחת התיק שייך לתוכנית אחרת עם אותו מספר — לא משייכים
        notes.append(f'בארכיון נמצא מסמך תב/{key} אחר ({(f or o).get("projectName", "")}, ועדה {committee:%d/%m/%Y}) — לא שויך')
        f, o, committee = {}, {}, None
    if re.sub(r'\D', '', planNum) != str(tik):
        notes.append('מס׳ התיק שונה ממס׳ התכנית')
    if sum(1 for x in RAW if x[0] == tik) > 1:
        notes.append('מס׳ תיק כפול')
    if not opened:
        notes.append('חסר תאריך פתיחה')
    rows.append([
        tik, planNum.replace('/ ', '/'), district, location, f.get('projectName') or o.get('projectName') or '',
        '' if change in ('', '0') else change, f.get('designation', ''), f.get('parcel', ''),
        f.get('developer', ''), f.get('devPhone', ''), f.get('devEmail', ''),
        f.get('archName') or ('' if submitter == '0' else submitter), f.get('archPhone', ''), f.get('archEmail', ''),
        f.get('contactName', ''), f.get('contactPhone', ''),
        opened, committee, o.get('session', ''), (o.get('decision') or '').strip(),
        status, handler, f.get('reliefsAsked', ''), '; '.join(notes),
    ])
rows.sort(key=lambda x: (x[16] or datetime.date.max, x[0]))

HEAD = ['מס׳ תיק', 'מס׳ תכנית', 'רובע', 'מיקום / מגרש', 'שם הפרויקט', 'תב״ע חלה', 'ייעוד', 'גוש / חלקה',
        'יזם', 'טלפון יזם', 'מייל יזם', 'אדריכל / מגיש', 'טלפון אדריכל', 'מייל אדריכל',
        'איש קשר', 'טלפון איש קשר', 'תאריך פתיחה', 'תאריך ועדה', 'מס׳ ישיבה', 'החלטת ועדה',
        'סטטוס', 'מטפלת', 'הקלות מבוקשות', 'הערות לבדיקה']
WIDTH = [9, 11, 11, 26, 34, 15, 22, 18, 28, 13, 26, 22, 13, 26, 16, 13, 12, 12, 10, 40, 11, 14, 50, 34]

BORD = PatternFill('solid', fgColor='7A1F2B')
thin = Side(style='thin', color='D9D4CC')
wb = Workbook()
ws = wb.active
ws.title = 'תוכניות בינוי'
ws.sheet_view.rightToLeft = True
ws.append(HEAD)
for r in rows: ws.append(r)
n = len(rows) + 1
for i, w in enumerate(WIDTH, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF', name='Arial')
    c.fill = BORD
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[1].height = 32
for row in ws.iter_rows(min_row=2, max_row=n):
    for c in row:
        c.font = Font(name='Arial', size=10)
        c.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True)
        c.border = Border(bottom=thin)
    for c in (row[16], row[17]):
        c.number_format = 'dd/mm/yyyy'
        c.alignment = Alignment(horizontal='center', vertical='top')
    if row[23].value:
        row[23].font = Font(name='Arial', size=10, color='9C5700')
tab = Table(displayName='Plans', ref=f'A1:{get_column_letter(len(HEAD))}{n}')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
ws.add_table(tab)
ws.freeze_panes = 'C2'

# ── גיליון סיכום (נוסחאות חיות על הטבלה) ──
s = wb.create_sheet('סיכום')
s.sheet_view.rightToLeft = True
s.column_dimensions['A'].width = 22; s.column_dimensions['B'].width = 14
def rng(col): return f"'תוכניות בינוי'!${col}$2:${col}${n}"
def block(title, label, keys, col):
    s.append([title]); s.cell(s.max_row, 1).font = Font(bold=True, size=12, name='Arial', color='7A1F2B')
    s.append([label, 'מס׳ תוכניות'])
    for c in s[s.max_row]: c.font = Font(bold=True, name='Arial')
    for k in keys:
        if col == 'Q':
            s.append([k, f'=COUNTIFS({rng("Q")},">="&DATE({k},1,1),{rng("Q")},"<"&DATE({k + 1},1,1))'])
        else:
            q = k.replace('"', '""'); s.append([k, f'=COUNTIF({rng(col)},"{q}")'])
    s.append([])
s.append(['סה״כ תוכניות', f'=COUNTA({rng("A")})']); s['A1'].font = Font(bold=True, name='Arial')
s.append(['עם תאריך ועדה', f'=COUNT({rng("R")})'])
s.append([])
block('לפי סטטוס', 'סטטוס', sorted({r[20] for r in rows}), 'U')
block('לפי שנת פתיחה', 'שנה', sorted({r[16].year for r in rows if r[16]}), 'Q')
block('לפי רובע', 'רובע', sorted({r[2] for r in rows}), 'C')

out = ROOT / 'exports' / 'plans.xlsx'
wb.save(out)
print(out, len(rows), 'rows')
